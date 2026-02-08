"""
OlingoBot — Telegram quiz bot.

Quiz structure: 5 blocks, 10 questions each.
Admin sends questions to all registered participants.
Answers and timestamps are saved to an Excel file.
"""

import json
import logging
import asyncio
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from telegram import (
    Update,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
)
from telegram.ext import (
    Application,
    CommandHandler,
    CallbackQueryHandler,
    ContextTypes,
)

from config import (
    BOT_TOKEN,
    ADMIN_IDS,
    QUESTION_TIME_LIMIT,
    EXCEL_FILE,
    QUESTIONS_FILE,
)

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Data structures
# ---------------------------------------------------------------------------

# Loaded question blocks from JSON
blocks: list[dict] = []

# Registered participants: {user_id: {"name": str, "username": str}}
participants: dict[int, dict] = {}

# Collected answers: list of dicts ready to be written to Excel
# Each entry: {user_id, user_name, username, block, question_id, question_text,
#              selected_option, is_correct, timestamp}
answers: list[dict] = []

# Track which question is currently active so we can reject late answers
# Key: question_id, Value: True while the question is open
active_questions: dict[int, bool] = {}

# Track which users already answered the current question
answered_users: dict[int, set[int]] = {}  # question_id -> set of user_ids


def load_questions() -> None:
    """Load questions from the JSON file."""
    global blocks
    path = Path(QUESTIONS_FILE)
    if not path.exists():
        logger.error("Questions file %s not found!", QUESTIONS_FILE)
        return
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    blocks = data.get("blocks", [])
    total = sum(len(b["questions"]) for b in blocks)
    logger.info("Loaded %d blocks, %d questions total.", len(blocks), total)


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------

EXCEL_HEADERS = [
    "User ID",
    "Имя",
    "Username",
    "Блок",
    "Вопрос №",
    "Текст вопроса",
    "Выбранный ответ",
    "Правильный?",
    "Время ответа (UTC)",
]


def save_to_excel() -> None:
    """Save all collected answers to the Excel file."""
    path = Path(EXCEL_FILE)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ответы"

    # Headers
    for col_idx, header in enumerate(EXCEL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = cell.font.copy(bold=True)

    # Data rows
    for row_idx, entry in enumerate(answers, start=2):
        ws.cell(row=row_idx, column=1, value=entry["user_id"])
        ws.cell(row=row_idx, column=2, value=entry["user_name"])
        ws.cell(row=row_idx, column=3, value=entry["username"])
        ws.cell(row=row_idx, column=4, value=entry["block"])
        ws.cell(row=row_idx, column=5, value=entry["question_id"])
        ws.cell(row=row_idx, column=6, value=entry["question_text"])
        ws.cell(row=row_idx, column=7, value=entry["selected_option"])
        ws.cell(row=row_idx, column=8, value="Да" if entry["is_correct"] else "Нет")
        ws.cell(row=row_idx, column=9, value=entry["timestamp"])

    # Auto-width columns
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = str(cell.value) if cell.value else ""
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max_len + 3, 60)

    wb.save(path)
    logger.info("Results saved to %s (%d rows).", EXCEL_FILE, len(answers))


# ---------------------------------------------------------------------------
# Access helpers
# ---------------------------------------------------------------------------


def is_admin(user_id: int) -> bool:
    return user_id in ADMIN_IDS


# ---------------------------------------------------------------------------
# Command handlers
# ---------------------------------------------------------------------------


async def cmd_start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Register participant or greet admin."""
    user = update.effective_user
    if is_admin(user.id):
        await update.message.reply_text(
            "Привет, админ! Команды:\n"
            "/send_block <номер> — отправить блок вопросов (1-5)\n"
            "/send_question <блок> <вопрос> — отправить один вопрос\n"
            "/participants — список участников\n"
            "/results — скачать Excel с результатами\n"
            "/standings — текущая таблица результатов"
        )
        return

    participants[user.id] = {
        "name": user.full_name,
        "username": user.username or "",
    }
    await update.message.reply_text(
        f"Добро пожаловать в викторину, {user.first_name}!\n"
        "Ожидайте — ведущий скоро начнёт отправлять вопросы."
    )
    logger.info("Participant registered: %s (id=%d)", user.full_name, user.id)


async def cmd_participants(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Show the list of registered participants (admin only)."""
    if not is_admin(update.effective_user.id):
        await update.message.reply_text("Эта команда доступна только администратору.")
        return

    if not participants:
        await update.message.reply_text("Пока нет зарегистрированных участников.")
        return

    lines = []
    for uid, info in participants.items():
        uname = f"@{info['username']}" if info["username"] else "—"
        lines.append(f"• {info['name']} ({uname}) [ID: {uid}]")

    await update.message.reply_text(
        f"Участники ({len(participants)}):\n" + "\n".join(lines)
    )


# ---------------------------------------------------------------------------
# Sending questions
# ---------------------------------------------------------------------------

OPTION_LABELS = ["A", "B", "C", "D"]


def build_question_message(
    block_name: str, q: dict
) -> tuple[str, InlineKeyboardMarkup]:
    """Build message text and inline keyboard for a question."""
    text_lines = [
        f"📦 *{block_name}*",
        f"❓ Вопрос {q['id']}: {q['text']}",
        "",
    ]
    buttons = []
    for idx, option in enumerate(q["options"]):
        text_lines.append(f"  {OPTION_LABELS[idx]}. {option}")
        buttons.append(
            InlineKeyboardButton(
                text=OPTION_LABELS[idx],
                callback_data=f"ans_{q['id']}_{idx}",
            )
        )

    if QUESTION_TIME_LIMIT > 0:
        text_lines.append(f"\n⏱ Время на ответ: {QUESTION_TIME_LIMIT} сек.")

    keyboard = InlineKeyboardMarkup([buttons])
    return "\n".join(text_lines), keyboard


async def send_question_to_all(
    block_name: str, q: dict, context: ContextTypes.DEFAULT_TYPE
) -> int:
    """Send a question to every registered participant. Returns count sent."""
    text, keyboard = build_question_message(block_name, q)
    active_questions[q["id"]] = True
    answered_users[q["id"]] = set()

    sent = 0
    for uid in participants:
        try:
            await context.bot.send_message(
                chat_id=uid,
                text=text,
                reply_markup=keyboard,
                parse_mode="Markdown",
            )
            sent += 1
        except Exception as e:
            logger.warning("Failed to send to %d: %s", uid, e)

    # Schedule closing the question after the time limit
    if QUESTION_TIME_LIMIT > 0:
        asyncio.get_event_loop().call_later(
            QUESTION_TIME_LIMIT,
            lambda qid=q["id"]: asyncio.ensure_future(
                close_question(qid, context)
            ),
        )

    return sent


async def close_question(question_id: int, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Mark question as closed after time limit."""
    if question_id in active_questions:
        active_questions[question_id] = False
        logger.info("Question %d closed (time limit).", question_id)


async def cmd_send_block(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Send an entire block of questions one by one (admin only)."""
    if not is_admin(update.effective_user.id):
        await update.message.reply_text("Эта команда доступна только администратору.")
        return

    if not context.args or not context.args[0].isdigit():
        await update.message.reply_text("Использование: /send_block <номер блока 1-5>")
        return

    block_num = int(context.args[0])
    if block_num < 1 or block_num > len(blocks):
        await update.message.reply_text(
            f"Блок {block_num} не найден. Доступно блоков: {len(blocks)}."
        )
        return

    if not participants:
        await update.message.reply_text(
            "Нет зарегистрированных участников. Участники должны сначала написать /start боту."
        )
        return

    block = blocks[block_num - 1]
    block_name = block["name"]
    questions = block["questions"]

    await update.message.reply_text(
        f"Начинаю отправку блока «{block_name}» ({len(questions)} вопросов)..."
    )

    for i, q in enumerate(questions):
        sent = await send_question_to_all(block_name, q, context)
        await update.message.reply_text(
            f"Вопрос {q['id']} отправлен {sent} участникам."
        )

        # Wait for the time limit + small buffer before sending next question
        if QUESTION_TIME_LIMIT > 0 and i < len(questions) - 1:
            await asyncio.sleep(QUESTION_TIME_LIMIT + 3)

    save_to_excel()
    await update.message.reply_text(
        f"Блок «{block_name}» завершён. Результаты сохранены в {EXCEL_FILE}."
    )


async def cmd_send_question(
    update: Update, context: ContextTypes.DEFAULT_TYPE
) -> None:
    """Send a single question (admin only). Usage: /send_question <block> <question>"""
    if not is_admin(update.effective_user.id):
        await update.message.reply_text("Эта команда доступна только администратору.")
        return

    if len(context.args) < 2:
        await update.message.reply_text(
            "Использование: /send_question <номер блока> <номер вопроса>\n"
            "Пример: /send_question 1 3"
        )
        return

    try:
        block_num = int(context.args[0])
        q_num = int(context.args[1])
    except ValueError:
        await update.message.reply_text("Номер блока и вопроса должны быть числами.")
        return

    if block_num < 1 or block_num > len(blocks):
        await update.message.reply_text(f"Блок {block_num} не найден.")
        return

    block = blocks[block_num - 1]
    if q_num < 1 or q_num > len(block["questions"]):
        await update.message.reply_text(
            f"Вопрос {q_num} не найден в блоке {block_num}."
        )
        return

    if not participants:
        await update.message.reply_text("Нет зарегистрированных участников.")
        return

    q = block["questions"][q_num - 1]
    sent = await send_question_to_all(block["name"], q, context)
    await update.message.reply_text(f"Вопрос {q['id']} отправлен {sent} участникам.")


# ---------------------------------------------------------------------------
# Answer handling
# ---------------------------------------------------------------------------


async def handle_answer(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Process an inline button press (participant's answer)."""
    query = update.callback_query
    await query.answer()

    user = query.from_user
    data = query.data  # format: ans_{question_id}_{option_index}

    parts = data.split("_")
    if len(parts) != 3 or parts[0] != "ans":
        return

    try:
        question_id = int(parts[1])
        option_idx = int(parts[2])
    except ValueError:
        return

    # Check if question is still active
    if not active_questions.get(question_id, False):
        await query.edit_message_reply_markup(reply_markup=None)
        await query.message.reply_text("⏰ Время на этот вопрос истекло.")
        return

    # Check if user already answered this question
    if user.id in answered_users.get(question_id, set()):
        await query.message.reply_text("Вы уже ответили на этот вопрос.")
        return

    answered_users.setdefault(question_id, set()).add(user.id)

    # Find the question
    q_data = None
    block_name = ""
    for block in blocks:
        for q in block["questions"]:
            if q["id"] == question_id:
                q_data = q
                block_name = block["name"]
                break
        if q_data:
            break

    if not q_data:
        return

    is_correct = option_idx == q_data["answer"]
    selected_text = q_data["options"][option_idx] if option_idx < len(q_data["options"]) else "?"
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")

    # Save answer
    answers.append(
        {
            "user_id": user.id,
            "user_name": user.full_name,
            "username": user.username or "",
            "block": block_name,
            "question_id": question_id,
            "question_text": q_data["text"],
            "selected_option": f"{OPTION_LABELS[option_idx]}. {selected_text}",
            "is_correct": is_correct,
            "timestamp": timestamp,
        }
    )

    # Update participant info if not registered
    if user.id not in participants:
        participants[user.id] = {
            "name": user.full_name,
            "username": user.username or "",
        }

    # Remove keyboard and confirm
    await query.edit_message_reply_markup(reply_markup=None)

    if is_correct:
        await query.message.reply_text("✅ Правильно!")
    else:
        correct_text = q_data["options"][q_data["answer"]]
        correct_label = OPTION_LABELS[q_data["answer"]]
        await query.message.reply_text(
            f"❌ Неправильно. Верный ответ: {correct_label}. {correct_text}"
        )

    logger.info(
        "Answer from %s (id=%d): Q%d -> %s (%s)",
        user.full_name,
        user.id,
        question_id,
        OPTION_LABELS[option_idx],
        "correct" if is_correct else "wrong",
    )


# ---------------------------------------------------------------------------
# Results
# ---------------------------------------------------------------------------


async def cmd_results(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Send the Excel results file to the admin."""
    if not is_admin(update.effective_user.id):
        await update.message.reply_text("Эта команда доступна только администратору.")
        return

    if not answers:
        await update.message.reply_text("Пока нет собранных ответов.")
        return

    save_to_excel()
    path = Path(EXCEL_FILE)
    if not path.exists():
        await update.message.reply_text("Файл результатов не найден.")
        return

    await update.message.reply_document(
        document=open(path, "rb"),
        filename=EXCEL_FILE,
        caption=f"Результаты викторины ({len(answers)} ответов)",
    )


async def cmd_standings(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Show current standings / leaderboard (admin only)."""
    if not is_admin(update.effective_user.id):
        await update.message.reply_text("Эта команда доступна только администратору.")
        return

    if not answers:
        await update.message.reply_text("Пока нет собранных ответов.")
        return

    # Calculate scores
    scores: dict[int, dict] = {}
    for entry in answers:
        uid = entry["user_id"]
        if uid not in scores:
            scores[uid] = {"name": entry["user_name"], "correct": 0, "total": 0}
        scores[uid]["total"] += 1
        if entry["is_correct"]:
            scores[uid]["correct"] += 1

    # Sort by correct answers descending
    sorted_scores = sorted(scores.values(), key=lambda x: x["correct"], reverse=True)

    lines = ["🏆 *Таблица результатов:*\n"]
    for i, s in enumerate(sorted_scores, start=1):
        lines.append(f"{i}. {s['name']} — {s['correct']}/{s['total']}")

    await update.message.reply_text("\n".join(lines), parse_mode="Markdown")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def main() -> None:
    load_questions()

    if BOT_TOKEN == "YOUR_BOT_TOKEN_HERE":
        logger.error(
            "Bot token not configured! Set BOT_TOKEN environment variable "
            "or edit config.py."
        )
        return

    if not ADMIN_IDS:
        logger.warning(
            "No admin IDs configured. Set ADMIN_IDS environment variable "
            "(comma-separated Telegram user IDs)."
        )

    app = Application.builder().token(BOT_TOKEN).build()

    app.add_handler(CommandHandler("start", cmd_start))
    app.add_handler(CommandHandler("participants", cmd_participants))
    app.add_handler(CommandHandler("send_block", cmd_send_block))
    app.add_handler(CommandHandler("send_question", cmd_send_question))
    app.add_handler(CommandHandler("results", cmd_results))
    app.add_handler(CommandHandler("standings", cmd_standings))
    app.add_handler(CallbackQueryHandler(handle_answer, pattern=r"^ans_"))

    logger.info("Bot started. Admin IDs: %s", ADMIN_IDS)
    app.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
